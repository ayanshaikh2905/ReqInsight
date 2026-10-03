import io
import os
from typing import Tuple, List
import docx
import pymupdf as fitz


class DocumentParser:
    """
    Extracts text from PDF, DOCX, and TXT files.
    Preserves structural section cues where possible.
    """

    @staticmethod
    def parse_pdf(file_bytes: bytes) -> str:
        """Extract text from PDF using PyMuPDF (fitz)."""
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text_parts: List[str] = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text.strip():
                text_parts.append(f"\n--- Page {page_num + 1} ---\n{page_text}")
        doc.close()
        return "\n".join(text_parts)

    @staticmethod
    def parse_docx(file_bytes: bytes) -> str:
        """Extract text from DOCX using python-docx."""
        file_stream = io.BytesIO(file_bytes)
        doc = docx.Document(file_stream)
        text_parts: List[str] = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Also parse tables if any
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_parts.append(" | ".join(row_texts))

        return "\n".join(text_parts)

    @staticmethod
    def parse_txt(file_bytes: bytes) -> str:
        """Extract text from TXT with multiple encoding fallbacks."""
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="replace")

    @classmethod
    def extract_text(cls, filename: str, file_bytes: bytes) -> str:
        """Dispatches parser based on file extension."""
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            return cls.parse_pdf(file_bytes)
        elif ext in [".docx", ".doc"]:
            return cls.parse_docx(file_bytes)
        elif ext in [".txt", ".md", ".csv", ".json"]:
            return cls.parse_txt(file_bytes)
        else:
            raise ValueError(f"Unsupported file format '{ext}'. Please upload a PDF, DOCX, or TXT file.")
